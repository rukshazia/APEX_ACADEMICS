import io
import csv
import re
from typing import List, Dict, Any, Tuple, Optional
import openpyxl

CANONICAL_FIELDS = {
    "roll_number": ["roll_number", "roll no", "roll_no", "reg_no", "register_no", "student_id", "id", "admission_no"],
    "full_name": ["full_name", "name", "student_name", "first_name", "candidate_name"],
    "email": ["email", "student_email", "mail_id", "email_address", "contact_email"],
    "department_code": ["department", "dept", "department_code", "dept_code", "branch", "program"],
    "semester": ["semester", "sem", "term", "current_semester"],
    "section": ["section", "sec", "division", "batch"],
    "cgpa": ["cgpa", "gpa", "percentage", "score", "grade_point"],
    "phone": ["phone", "mobile", "contact_no", "phone_number"],
    "guardian_name": ["guardian_name", "parent_name", "father_name", "guardian"],
    "guardian_phone": ["guardian_phone", "parent_phone", "parent_mobile"],
    "placement_status": ["placement_status", "placement", "placement_readiness", "status"]
}

def guess_column_mapping(headers: List[str]) -> Dict[str, str]:
    """Suggests mapping from file headers to canonical student fields."""
    mapping: Dict[str, str] = {}
    normalized_headers = [h.strip().lower().replace("-", "_").replace(" ", "_") for h in headers]

    for canonical, synonyms in CANONICAL_FIELDS.items():
        found = False
        for syn in synonyms:
            syn_norm = syn.strip().lower().replace("-", "_").replace(" ", "_")
            for idx, orig_header in enumerate(headers):
                nh = normalized_headers[idx]
                if nh == syn_norm or syn_norm in nh:
                    mapping[canonical] = orig_header
                    found = True
                    break
            if found:
                break
    return mapping

def inspect_workbook(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Inspects sheets and extracts headers and first 5 rows for preview."""
    is_csv = filename.lower().endswith(".csv")
    sheets_info = []

    if is_csv:
        content = file_bytes.decode('utf-8', errors='replace')
        csv_reader = csv.reader(io.StringIO(content))
        rows = list(csv_reader)
        if not rows:
            return {"sheets": [], "error": "File is empty"}
        headers = [str(c).strip() for c in rows[0] if str(c).strip() != ""]
        sample_rows = rows[1:6]
        sheets_info.append({
            "name": "Sheet1",
            "headers": headers,
            "sample_rows": sample_rows,
            "total_rows_estimated": len(rows) - 1
        })
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        for name in sheet_names:
            ws = wb[name]
            rows_iter = ws.iter_rows(values_only=True)
            try:
                first_row = next(rows_iter, None)
            except Exception:
                first_row = None
            if not first_row:
                continue
            headers = [str(c).strip() for c in first_row if c is not None and str(c).strip() != ""]
            sample_rows = []
            for _ in range(5):
                r = next(rows_iter, None)
                if r is None:
                    break
                sample_rows.append([str(c) if c is not None else "" for c in r[:len(headers)]])
            
            # Count total rows roughly
            estimated_count = ws.max_row - 1 if ws.max_row else len(sample_rows)
            sheets_info.append({
                "name": name,
                "headers": headers,
                "sample_rows": sample_rows,
                "total_rows_estimated": max(0, estimated_count)
            })

    if not sheets_info:
        return {"sheets": [], "error": "No valid data sheets found in workbook"}

    first_sheet = sheets_info[0]
    suggested_mapping = guess_column_mapping(first_sheet["headers"])

    return {
        "filename": filename,
        "sheets": sheets_info,
        "suggested_mapping": suggested_mapping,
        "canonical_fields": list(CANONICAL_FIELDS.keys())
    }

def parse_and_validate_records(
    file_bytes: bytes,
    filename: str,
    sheet_name: Optional[str],
    column_mapping: Dict[str, str], # canonical -> source_header
    existing_rolls: set,
    existing_emails: set
) -> Dict[str, Any]:
    """Parses workbook sheet with column mapping, validates rules, detects duplicates & conflicts."""
    is_csv = filename.lower().endswith(".csv")
    raw_rows: List[Dict[str, Any]] = []

    if is_csv:
        content = file_bytes.decode('utf-8', errors='replace')
        reader = csv.DictReader(io.StringIO(content))
        for r in reader:
            raw_rows.append(r)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        target_sheet = sheet_name if sheet_name and sheet_name in wb.sheetnames else wb.sheetnames[0]
        ws = wb[target_sheet]
        rows_vals = list(ws.iter_rows(values_only=True))
        if not rows_vals:
            return {"error": "Sheet is empty"}
        headers = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(rows_vals[0])]
        for r in rows_vals[1:]:
            if not any(r): # skip blank rows
                continue
            row_dict = {}
            for idx, h in enumerate(headers):
                val = r[idx] if idx < len(r) else ""
                row_dict[h] = str(val).strip() if val is not None else ""
            raw_rows.append(row_dict)

    # Inverted mapping: source_header -> canonical
    inv_map = {src_col: canon for canon, src_col in column_mapping.items() if src_col}

    seen_rolls = set()
    seen_emails = set()

    to_create: List[Dict[str, Any]] = []
    to_update: List[Dict[str, Any]] = []
    rejected: List[Dict[str, Any]] = []

    EMAIL_REGEX = re.compile(r"^[\w\.-]+@[\w\.-]+\.\w+$")

    for row_idx, r in enumerate(raw_rows, start=2): # 1-indexed row number, 1 is header
        clean_item = {}
        for src_col, val in r.items():
            if src_col in inv_map:
                clean_item[inv_map[src_col]] = str(val).strip()

        row_errors: List[str] = []

        # Validate roll_number
        roll = clean_item.get("roll_number", "").strip()
        if not roll:
            row_errors.append("Missing required Roll Number / Student ID")
        elif roll in seen_rolls:
            row_errors.append(f"Duplicate Roll Number '{roll}' within the file")
        else:
            seen_rolls.add(roll)

        # Validate email
        email = clean_item.get("email", "").strip().lower()
        if not email:
            row_errors.append("Missing required Email address")
        elif not EMAIL_REGEX.match(email):
            row_errors.append(f"Invalid email format '{email}'")
        elif email in seen_emails:
            row_errors.append(f"Duplicate Email '{email}' within the file")
        else:
            seen_emails.add(email)

        # Validate full_name
        name = clean_item.get("full_name", "").strip()
        if not name:
            row_errors.append("Missing required Student Name")

        # Validate semester
        sem_val = clean_item.get("semester", "1")
        try:
            sem_int = int(float(sem_val)) if sem_val else 1
            if sem_int < 1 or sem_int > 12:
                row_errors.append(f"Semester '{sem_val}' must be between 1 and 12")
            clean_item["semester"] = sem_int
        except ValueError:
            row_errors.append(f"Semester '{sem_val}' is not a valid number")
            clean_item["semester"] = 1

        # Validate CGPA
        cgpa_val = clean_item.get("cgpa", "0.0")
        try:
            cgpa_float = float(cgpa_val) if cgpa_val else 0.0
            if cgpa_float < 0.0 or cgpa_float > 10.0:
                row_errors.append(f"CGPA '{cgpa_val}' must be between 0.0 and 10.0")
            clean_item["cgpa"] = cgpa_float
        except ValueError:
            row_errors.append(f"CGPA '{cgpa_val}' is not a valid number")
            clean_item["cgpa"] = 0.0

        if row_errors:
            rejected.append({
                "row_number": row_idx,
                "raw_data": clean_item,
                "errors": row_errors
            })
        else:
            # Check if student exists in database (conflict/update handling)
            is_conflict = (roll in existing_rolls) or (email in existing_emails)
            clean_item["row_number"] = row_idx
            if is_conflict:
                to_update.append(clean_item)
            else:
                to_create.append(clean_item)

    return {
        "total_rows_processed": len(raw_rows),
        "create_count": len(to_create),
        "update_count": len(to_update),
        "rejected_count": len(rejected),
        "to_create_preview": to_create[:20],
        "to_update_preview": to_update[:20],
        "rejected_rows": rejected,
        "valid_to_create": to_create,
        "valid_to_update": to_update
    }
